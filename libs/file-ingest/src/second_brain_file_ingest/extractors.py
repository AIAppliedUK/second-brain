from __future__ import annotations

import hashlib
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import xlrd
from pypdf import PdfReader
from pypdf.errors import PdfReadError, PdfStreamError


class FileExtractionError(RuntimeError):
    pass


@dataclass(slots=True)
class ExtractedContent:
    text: str
    metadata: dict[str, str]
    content_hash: str


class FileExtractor:
    plain_text_suffixes = {".md", ".txt"}
    supported_suffixes = plain_text_suffixes | {".pdf", ".docx", ".xlsx", ".xls", ".archimate"}

    def extract(self, path: Path) -> ExtractedContent:
        suffix = path.suffix.lower()
        if suffix not in self.supported_suffixes:
            raise FileExtractionError(f"Unsupported file type: {path.suffix}")
        raw_bytes = path.read_bytes()
        content_hash = hashlib.sha256(raw_bytes).hexdigest()
        if suffix in self.plain_text_suffixes:
            text = path.read_text(encoding="utf-8", errors="replace")
        elif suffix == ".pdf":
            try:
                reader = PdfReader(str(path))
                text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
            except (PdfReadError, PdfStreamError, OSError, ValueError) as exc:
                raise FileExtractionError(f"Failed to extract PDF {path}: {exc}") from exc
        elif suffix == ".docx":
            try:
                text = self._extract_docx(path)
            except (KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
                raise FileExtractionError(f"Failed to extract DOCX {path}: {exc}") from exc
        elif suffix == ".xlsx":
            try:
                text = self._extract_xlsx(path)
            except (KeyError, OSError, ValueError, zipfile.BadZipFile) as exc:
                raise FileExtractionError(f"Failed to extract XLSX {path}: {exc}") from exc
        elif suffix == ".xls":
            try:
                text = self._extract_xls(path)
            except FileExtractionError:
                raise
            except Exception as exc:
                raise FileExtractionError(f"Failed to extract XLS {path}: {exc}") from exc
        elif suffix == ".archimate":
            try:
                text = self._extract_archimate(path)
            except FileExtractionError:
                raise
            except Exception as exc:
                raise FileExtractionError(f"Failed to extract ArchiMate {path}: {exc}") from exc
        else:
            raise FileExtractionError(f"Unsupported file type: {path.suffix}")
        text = self._sanitize_text(text)
        if not text.strip():
            raise FileExtractionError(f"No extractable text found in {path}")
        return ExtractedContent(
            text=text.strip(),
            metadata={"suffix": suffix, "size_bytes": str(path.stat().st_size)},
            content_hash=content_hash,
        )

    def _sanitize_text(self, text: str) -> str:
        # Some PDF/doc exports contain NUL bytes, which PostgreSQL text fields reject.
        return text.replace("\x00", "")

    def _extract_docx(self, path: Path) -> str:
        with zipfile.ZipFile(path) as archive:
            document = archive.read("word/document.xml").decode("utf-8", errors="ignore")
        fragments: list[str] = []
        for segment in document.replace("</w:p>", "\n").split("<"):
            if segment.startswith("w:t") and ">" in segment:
                fragments.append(segment.split(">", 1)[1])
        return "".join(fragments).replace("&amp;", "&").strip()

    def _extract_xls(self, path: Path) -> str:
        wb = xlrd.open_workbook(str(path))
        sections: list[str] = []
        for sheet_index in range(wb.nsheets):
            ws = wb.sheet_by_index(sheet_index)
            if ws.nrows == 0:
                continue
            rows: list[str] = []
            for row_index in range(ws.nrows):
                cells: list[str] = []
                for col_index in range(ws.ncols):
                    cell = ws.cell(row_index, col_index)
                    if cell.ctype == xlrd.XL_CELL_NUMBER:
                        value = cell.value
                        cells.append(str(int(value)) if value == int(value) else str(value))
                    else:
                        cells.append(str(cell.value))
                rows.append(" | ".join(cells))
            if not rows:
                continue
            lines: list[str] = [f"## Sheet: {ws.name}"] + rows
            sections.append("\n".join(lines))
        return "\n\n".join(sections)

    def _extract_archimate(self, path: Path) -> str:
        xsi_ns = "http://www.w3.org/2001/XMLSchema-instance"

        layer_types: dict[str, str] = {}
        for t in [
            "BusinessActor",
            "BusinessRole",
            "BusinessCollaboration",
            "BusinessInterface",
            "BusinessProcess",
            "BusinessFunction",
            "BusinessInteraction",
            "BusinessEvent",
            "BusinessService",
            "BusinessObject",
            "Contract",
            "Representation",
            "Product",
        ]:
            layer_types[t] = "Business"
        for t in [
            "ApplicationComponent",
            "ApplicationCollaboration",
            "ApplicationInterface",
            "ApplicationFunction",
            "ApplicationInteraction",
            "ApplicationProcess",
            "ApplicationEvent",
            "ApplicationService",
            "DataObject",
        ]:
            layer_types[t] = "Application"
        for t in [
            "Node",
            "Device",
            "SystemSoftware",
            "TechnologyCollaboration",
            "TechnologyInterface",
            "Path",
            "CommunicationNetwork",
            "TechnologyFunction",
            "TechnologyProcess",
            "TechnologyInteraction",
            "TechnologyEvent",
            "TechnologyService",
            "Artifact",
        ]:
            layer_types[t] = "Technology"
        for t in ["Resource", "Capability", "CourseOfAction", "ValueStream"]:
            layer_types[t] = "Strategy"
        for t in [
            "Stakeholder",
            "Driver",
            "Assessment",
            "Goal",
            "Outcome",
            "Principle",
            "Requirement",
            "Constraint",
            "Meaning",
            "Value",
        ]:
            layer_types[t] = "Motivation"
        for t in ["WorkPackage", "Deliverable", "ImplementationEvent", "Plateau", "Gap"]:
            layer_types[t] = "Implementation"

        layer_order = [
            "Business",
            "Application",
            "Technology",
            "Strategy",
            "Motivation",
            "Implementation",
            "Other",
        ]

        tree = ET.parse(path)
        root = tree.getroot()

        # Collect all elements recursively
        id_to_name: dict[str, str] = {}
        elements: list[dict] = []
        relationships: list[dict] = []

        for elem in root.iter():
            xsi_type = elem.get(f"{{{xsi_ns}}}type")
            if xsi_type is None:
                continue
            # Strip namespace prefix (e.g. "archimate:BusinessProcess" -> "BusinessProcess")
            type_name = xsi_type.split(":", 1)[1] if ":" in xsi_type else xsi_type

            elem_id = elem.get("id")
            elem_name = elem.get("name", "")

            if elem_id and elem_name:
                id_to_name[elem_id] = elem_name

            if type_name.endswith("Relationship"):
                relationships.append(
                    {
                        "type": type_name,
                        "name": elem_name,
                        "source": elem.get("source", ""),
                        "target": elem.get("target", ""),
                    }
                )
            else:
                doc_elem = elem.find("documentation")
                doc_text = doc_elem.text.strip() if doc_elem is not None and doc_elem.text else ""
                elements.append(
                    {
                        "type": type_name,
                        "name": elem_name,
                        "doc": doc_text,
                        "layer": layer_types.get(type_name, "Other"),
                    }
                )

        def format_type(type_name: str) -> str:
            return re.sub(r"(?<=[a-z])(?=[A-Z])", " ", type_name)

        # Group elements by layer
        layers: dict[str, list[dict]] = {layer: [] for layer in layer_order}
        for elem in elements:
            layers[elem["layer"]].append(elem)

        sections: list[str] = []

        for layer in layer_order:
            layer_elements = layers[layer]
            if not layer_elements:
                continue
            lines: list[str] = [f"## {layer} Layer", ""]
            for elem in layer_elements:
                lines.append(f'{format_type(elem["type"])}: "{elem["name"]}"')
                if elem["doc"]:
                    lines.append(elem["doc"])
                lines.append("")
            # Remove trailing blank line
            while lines and lines[-1] == "":
                lines.pop()
            sections.append("\n".join(lines))

        if relationships:
            rel_lines: list[str] = ["## Relationships", ""]
            for rel in relationships:
                src_name = id_to_name.get(rel["source"], rel["source"])
                tgt_name = id_to_name.get(rel["target"], rel["target"])
                label = rel["name"] if rel["name"] else rel["type"]
                rel_lines.append(f'"{src_name}" -> "{tgt_name}" ({label})')
            sections.append("\n".join(rel_lines))

        return "\n\n".join(sections)

    def _extract_xlsx(self, path: Path) -> str:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sections: list[str] = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                # Filter to rows that have at least one non-None cell
                non_empty_rows = [row for row in rows if any(cell is not None for cell in row)]
                if not non_empty_rows:
                    continue
                lines: list[str] = [f"## Sheet: {sheet_name}"]
                for row in non_empty_rows:
                    cells: list[str] = []
                    for cell in row:
                        if cell is None:
                            cells.append("")
                        elif isinstance(cell, float) and cell == int(cell):
                            cells.append(str(int(cell)))
                        else:
                            cells.append(str(cell))
                    lines.append(" | ".join(cells))
                sections.append("\n".join(lines))
            return "\n\n".join(sections)
        finally:
            wb.close()
