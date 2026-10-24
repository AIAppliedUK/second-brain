"""Regression tests for the 2026-10 index-quality review.

Each test pins one finding from the review: failures retried forever, stale
sources never removed, empty-cell spreadsheet chunks, lexical hits starved out
of the fused ranking, the vector index bypassed, and the MCP default scope
hiding every other project from unscoped searches.
"""

from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

import openpyxl

from second_brain_core.config import MemoryRoot, Settings
from second_brain_core.embeddings import DeterministicEmbedder, build_embedder
from second_brain_core.repository import (
    build_lexical_filter,
    build_lexical_query,
    build_vector_query,
)
from second_brain_core.retrieval import HybridRetriever
from second_brain_file_ingest import FileExtractionError, FileExtractor, LocalFileIngester
from second_brain_ingester.cli import process_local_file, reconcile_removed_sources
from second_brain_mcp_server.service import MCPService
from second_brain_models import (
    ChunkRecord,
    SearchMemoryRequest,
    SearchRequest,
    SourceRecord,
)
from test_ingester_cli import FakeAuditRepository, FakeChunkRepository

# --- ingestion: failures ---------------------------------------------------


class FailingIngester:
    def __init__(self):
        self.extract_calls = 0

    def extract_path(self, path: Path):
        self.extract_calls += 1
        raise FileExtractionError(f"Failed to extract DOCX {path}: File is not a zip file")


class RecordingSourceRepository:
    def __init__(self, existing: SourceRecord | None = None):
        self.existing = existing
        self.failed = []

    def get_source_by_identity(self, memory_scope, source_type, external_id):
        return self.existing

    def mark_failed(self, source):
        self.failed.append(source)
        return SourceRecord(**source.model_dump())


def test_failed_file_is_recorded_so_unchanged_file_is_not_retried(tmp_path: Path):
    path = tmp_path / "locked.docx"
    path.write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 64)
    root = MemoryRoot(scope="project:test", path=tmp_path)
    ingester = FailingIngester()
    repo = RecordingSourceRepository()
    audit = FakeAuditRepository()

    assert process_local_file(root, path, ingester, repo, FakeChunkRepository(), audit) == "failed"
    assert len(repo.failed) == 1
    failed = repo.failed[0]
    assert failed.status == "failed"
    assert "not a zip file" in failed.metadata["error"]

    # Second pass: the recorded failure (same mtime + size) short-circuits extraction.
    repo.existing = SourceRecord(**failed.model_dump())
    assert process_local_file(root, path, ingester, repo, FakeChunkRepository(), audit) == "skipped"
    assert ingester.extract_calls == 1
    assert len(audit.events) == 1


def test_discovery_skips_office_lock_files_and_empty_files(tmp_path: Path):
    (tmp_path / "~$Schedule.docx").write_bytes(b"lock")
    (tmp_path / "empty.md").write_text("")
    (tmp_path / "real.md").write_text("# Real\n\ncontent")
    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    found = [p.name for _, p in ingester.discover_files([MemoryRoot("s", tmp_path)])]
    assert found == ["real.md"]


def test_discovery_skips_claude_agent_folders_at_any_depth(tmp_path: Path):
    for folder in (tmp_path / ".claude" / "sprint-state", tmp_path / "app" / ".claude"):
        folder.mkdir(parents=True)
        (folder / "notes.md").write_text("# Agent state\n\ncontent")
    (tmp_path / "real.md").write_text("# Real\n\ncontent")
    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    found = [p.name for _, p in ingester.discover_files([MemoryRoot("s", tmp_path)])]
    assert found == ["real.md"]


def test_encrypted_office_file_gets_a_clear_error(tmp_path: Path):
    path = tmp_path / "Q2.1.docx"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)
    try:
        FileExtractor().extract(path)
    except FileExtractionError as exc:
        assert "password-protected" in str(exc)
    else:
        raise AssertionError("expected FileExtractionError")


# --- ingestion: stale sources ----------------------------------------------


class ReconcileSourceRepository:
    def __init__(self, indexed: list[str]):
        self.indexed = indexed
        self.deleted: list[str] = []

    def list_local_file_ids(self, memory_scope):
        return {external_id: uuid4() for external_id in self.indexed}

    def delete_sources(self, source_ids):
        self.deleted.extend(source_ids)
        return len(source_ids)


def test_reconcile_removes_sources_whose_files_are_gone(tmp_path: Path):
    kept = tmp_path / "kept.md"
    kept.write_text("x")
    gone = str(tmp_path / ".claude" / "worktrees" / "agent-1" / "kept.md")
    repo = ReconcileSourceRepository([str(kept.resolve()), gone])
    root = MemoryRoot(scope="project:test", path=tmp_path)
    removed = reconcile_removed_sources(root, {kept.resolve()}, repo)
    assert removed == 1
    assert len(repo.deleted) == 1


def test_reconcile_does_nothing_when_root_is_unreadable(tmp_path: Path):
    repo = ReconcileSourceRepository(["/somewhere/file.md"])
    root = MemoryRoot(scope="project:test", path=tmp_path / "not-mounted")
    assert reconcile_removed_sources(root, set(), repo) == 0
    assert repo.deleted == []


# --- ingestion: content quality ---------------------------------------------


def test_xlsx_extraction_drops_empty_cells(tmp_path: Path):
    path = tmp_path / "pricing.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pricing"
    ws["A1"] = "Item"
    ws["F1"] = "Price"
    ws["A3"] = "Licence"
    ws["F3"] = 1200
    ws["Z40"] = "   "
    wb.save(path)
    text = FileExtractor().extract(path).text
    assert "|  |" not in text
    assert "Item | Price" in text
    assert "Licence | 1200" in text


def test_chunks_without_searchable_words_are_dropped(tmp_path: Path):
    path = tmp_path / "notes.md"
    path.write_text("# Notes\n\nReal content here.\n\n## Divider\n\n---\n\n| | | |\n")
    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    document, _ = ingester.ingest_path(path, memory_scope="s", root_path=tmp_path)
    assert [c.chunk_text for c in document.chunks] == ["Real content here."]


# --- retrieval ----------------------------------------------------------------


def _source() -> SourceRecord:
    return SourceRecord(
        id=uuid4(),
        source_type="local_file",
        external_id="f",
        title="t",
        uri="file:///t",
        content_hash="h",
        updated_at=datetime.now(UTC),
    )


def _chunk(source, **scores) -> ChunkRecord:
    return ChunkRecord(
        id=uuid4(),
        source_id=source.id,
        chunk_index=0,
        heading_path=[],
        chunk_text="x",
        token_count=1,
        embedding=[],
        source=source,
        **scores,
    )


class _Repo:
    def __init__(self, lexical, vector):
        self.lexical, self.vector = lexical, vector

    def lexical_search(self, _):
        return self.lexical

    def vector_search(self, _):
        return self.vector


def test_top_lexical_hit_is_not_starved_by_disjoint_vector_hits():
    source = _source()
    lexical = [_chunk(source, lexical_score=0.9)]
    vector = [_chunk(source, vector_score=0.5 - i / 100) for i in range(20)]
    retriever = HybridRetriever(_Repo(lexical, vector), DeterministicEmbedder(8))
    hits = retriever.search(SearchRequest(query="q", limit=10)).hits
    assert lexical[0].id in [h.chunk_id for h in hits[:2]]


def test_queries_only_return_active_sources():
    lexical = build_lexical_query("a b", 8, 5, {})
    vector = build_vector_query([0.1, 0.2], 5, {})
    assert "s.status = 'active'" in lexical.sql
    assert "s.status = 'active'" in vector.sql


def test_lexical_query_matches_any_term_not_all_terms():
    bundle = build_lexical_query("how do we handle disaster recovery", 8, 5, {})
    # websearch_to_tsquery ANDs every term; natural questions then match nothing.
    assert "websearch_to_tsquery" not in bundle.sql
    assert "'|'" in bundle.sql or "' | '" in bundle.sql


def test_lexical_filter_keeps_rare_terms_and_pairs_common_ones():
    # Common words ("chang" is in 11% of chunks) matched alone force a full table scan
    # for chunks that would rank last anyway; they only qualify a chunk in pairs.
    frequencies = {"chang": 0.11, "get": 0.05, "train": 0.018}
    tsq = build_lexical_filter(["train", "chang", "champion", "get"], frequencies)
    assert tsq == "'train' | 'champion' | ('chang' & 'get')"


def test_lexical_filter_falls_back_to_plain_or_without_statistics():
    assert build_lexical_filter(["audit", "record"], {}) == "'audit' | 'record'"
    assert build_lexical_filter(["chang"], {"chang": 0.5}) == "'chang'"
    assert build_lexical_filter([], {}) is None


def test_lexical_filter_quotes_lexemes():
    assert build_lexical_filter(["o'neil", "a\\b"], {}) == "'o''neil' | 'a\\\\b'"


def test_vector_query_orders_by_distance_only_so_the_ann_index_is_used():
    bundle = build_vector_query([0.1, 0.2], 5, {})
    order_by = bundle.sql.split("ORDER BY", 1)[1].split("LIMIT", 1)[0]
    assert "chunk_index" not in order_by


# --- MCP ------------------------------------------------------------------------


class _CapturingRetriever:
    def __init__(self):
        self.requests = []

    def search(self, request):
        self.requests.append(request)
        from second_brain_models import SearchResponse

        return SearchResponse(query=request.query, hits=[])


def test_unscoped_search_searches_every_scope_even_with_a_default_write_scope():
    retriever = _CapturingRetriever()
    service = MCPService(
        retriever, None, None, DeterministicEmbedder(8), default_memory_scope="project:x"
    )
    service.search_memory(SearchMemoryRequest(query="anything"))
    assert retriever.requests[0].filters.memory_scopes is None


# --- embeddings -------------------------------------------------------------------


def test_build_embedder_honours_deterministic_provider(monkeypatch):
    monkeypatch.setenv("SECOND_BRAIN_EMBEDDING_PROVIDER", "deterministic")
    monkeypatch.setenv("SECOND_BRAIN_EMBEDDING_DIMENSION", "16")
    embedder = build_embedder(Settings())
    assert isinstance(embedder, DeterministicEmbedder)
    assert embedder.dimension == 16
    assert embedder.embed_many(["a b", "c"]) == [embedder.embed("a b"), embedder.embed("c")]


def test_lexical_rank_puts_term_coverage_before_term_density():
    # With OR matching, ts_rank_cd alone ranks a chunk repeating one query word above a
    # chunk containing every query word. Coverage (distinct terms matched) must lead.
    bundle = build_lexical_query("schedule flexibility analysis", 8, 5, {})
    score_expr = bundle.sql.split("AS lexical_score", 1)[0].rsplit("c.metadata,", 1)[1]
    assert "unnest(q.lexemes)" in score_expr
    assert "ts_rank_cd(c.tsv, q.tsq, 32)" in score_expr


def _write_docx(path: Path, document_xml: str) -> None:
    import zipfile

    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", document_xml)


def test_docx_keeps_paragraph_breaks_and_decodes_entities(tmp_path: Path):
    path = tmp_path / "answer.docx"
    _write_docx(
        path,
        "<w:document><w:body>"
        "<w:p><w:r><w:t>Scriptable output</w:t></w:r></w:p>"
        '<w:p><w:r><w:t xml:space="preserve">Request &quot;rate&quot; </w:t></w:r>'
        "<w:r><w:tab/><w:t>limits &lt;300ms &amp; more</w:t></w:r></w:p>"
        "</w:body></w:document>",
    )
    text = FileExtractor().extract(path).text
    assert text == 'Scriptable output\nRequest "rate"  limits <300ms & more'


def test_file_indexed_by_an_older_extractor_is_reprocessed(tmp_path: Path):
    from second_brain_file_ingest.extractors import EXTRACTOR_VERSION

    path = tmp_path / "notes.md"
    path.write_text("hello")
    mtime = datetime.fromtimestamp(path.stat().st_mtime, tz=UTC)
    existing = SourceRecord(
        id=uuid4(),
        memory_scope="s",
        source_type="local_file",
        external_id=str(path.resolve()),
        title=path.name,
        uri=path.resolve().as_uri(),
        content_hash="same",
        metadata={"size_bytes": str(path.stat().st_size), "extractor_version": "0"},
        updated_at=mtime,
    )
    ingester = FailingIngester()
    process_local_file(
        MemoryRoot("s", tmp_path),
        path,
        ingester,
        RecordingSourceRepository(existing),
        FakeChunkRepository(),
        FakeAuditRepository(),
    )
    assert ingester.extract_calls == 1
    assert EXTRACTOR_VERSION != "0"


class _ReverseReranker:
    """Scores later candidates higher, so a reranked list comes back reversed."""

    def __init__(self):
        self.calls = []

    def score(self, query, passages):
        self.calls.append((query, list(passages)))
        return [float(i) for i in range(len(passages))]


def test_reranker_reorders_fused_candidates_and_trims_to_limit():
    source = _source()
    lexical = [_chunk(source, lexical_score=1.0 - i / 10) for i in range(4)]
    reranker = _ReverseReranker()
    retriever = HybridRetriever(_Repo(lexical, []), DeterministicEmbedder(8), reranker=reranker)
    hits = retriever.search(SearchRequest(query="q", limit=2, rerank=True)).hits
    assert [h.chunk_id for h in hits] == [lexical[3].id, lexical[2].id]
    assert reranker.calls[0][0] == "q"
    assert len(reranker.calls[0][1]) == 4  # every fused candidate is rescored, not just `limit`


def test_reranker_only_runs_for_deep_searches():
    # Reranking costs ~2s per search on CPU, so it is opt-in per request.
    source = _source()
    reranker = _ReverseReranker()
    retriever = HybridRetriever(
        _Repo([_chunk(source, lexical_score=1.0)], []), DeterministicEmbedder(8), reranker=reranker
    )
    retriever.search(SearchRequest(query="q", limit=2))
    assert reranker.calls == []


def test_deep_mcp_search_requests_reranking():
    retriever = _CapturingRetriever()
    service = MCPService(retriever, None, None, DeterministicEmbedder(8))
    service.search_memory(SearchMemoryRequest(query="q", deep=True))
    service.search_memory(SearchMemoryRequest(query="q"))
    assert [r.rerank for r in retriever.requests] == [True, False]


# --- whole-folder indexing and housekeeping ---------------------------------------


def test_every_project_subfolder_becomes_its_own_scope(tmp_path: Path, monkeypatch):
    (tmp_path / "alpha").mkdir()
    (tmp_path / "beta build").mkdir()
    (tmp_path / ".cache").mkdir()
    (tmp_path / "pat.txt").write_text("secret")  # loose top-level files are never roots
    monkeypatch.setenv("SECOND_BRAIN_PROJECT_PARENTS", str(tmp_path))
    monkeypatch.setenv("SECOND_BRAIN_MEMORY_ROOTS", f"project:alpha-explicit|{tmp_path / 'alpha'}")
    roots = {root.scope: root.path for root in Settings().resolved_memory_roots()}
    # The explicit root for alpha wins over the auto-discovered one; no duplicates.
    assert roots == {
        "project:alpha-explicit": tmp_path / "alpha",
        "project:beta build": tmp_path / "beta build",
    }


class FolderCleanupRepository:
    def __init__(self):
        self.calls = []

    def delete_local_files_outside_scopes(self, path_prefix, keep_scopes):
        self.calls.append((path_prefix, set(keep_scopes)))
        return 3


def test_deleted_project_folders_are_purged_from_the_index(tmp_path: Path):
    from second_brain_ingester.cli import remove_deleted_projects

    (tmp_path / "alpha").mkdir()
    roots = [MemoryRoot("project:alpha", tmp_path / "alpha"), MemoryRoot("onedrive", Path("/x"))]
    repo = FolderCleanupRepository()
    assert remove_deleted_projects([tmp_path], roots, repo) == 3
    assert repo.calls == [(str(tmp_path.resolve()), {"project:alpha"})]


def test_project_purge_is_skipped_when_the_parent_is_not_mounted(tmp_path: Path):
    from second_brain_ingester.cli import remove_deleted_projects

    repo = FolderCleanupRepository()
    assert remove_deleted_projects([tmp_path / "missing"], [], repo) == 0
    (tmp_path / "empty").mkdir()
    assert remove_deleted_projects([tmp_path / "empty"], [], repo) == 0
    assert repo.calls == []


def test_remember_refuses_to_overwrite_indexed_files_or_onenote_pages():
    # remember() upserts by (scope, source_type, external_id); with an ingester-owned
    # type a caller could replace a real file's indexed text with arbitrary content.
    import pytest

    from second_brain_models import RememberRequest

    service = MCPService(_CapturingRetriever(), None, None, DeterministicEmbedder(8))
    for source_type in ("local_file", "onenote"):
        request = RememberRequest(
            body="poisoned", source_type=source_type, external_id="/x/real.docx"
        )
        with pytest.raises(ValueError, match="source_type"):
            service.remember(request)
