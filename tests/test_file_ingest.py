from pathlib import Path

from second_brain_core.config import MemoryRoot
from second_brain_core.embeddings import DeterministicEmbedder
from second_brain_file_ingest import FileExtractionError, FileExtractor, LocalFileIngester


def test_local_file_ingester_builds_canonical_document_for_markdown(tmp_path: Path):
    path = tmp_path / "notes.md"
    path.write_text(
        "# Heading\n\nThis is a note about retrieval.\n\n## Details\n\nCitations matter."
    )
    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    document, event = ingester.ingest_path(
        path, memory_scope="project:test", root_path=tmp_path
    )
    assert document is not None
    assert document.source.memory_scope == "project:test"
    assert document.source.source_type == "local_file"
    assert len(document.chunks) >= 1
    assert document.chunks[0].memory_scope == "project:test"
    assert document.chunks[0].heading_path == ["Heading"]
    assert event.status == "success"


def test_local_file_ingester_logs_failures_without_document(tmp_path: Path):
    path = tmp_path / "notes.bin"
    path.write_bytes(b"\x00\x01")
    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    document, event = ingester.ingest_path(path, memory_scope="general", root_path=tmp_path)
    assert document is None
    assert event.status == "failed"


def test_discover_files_returns_scoped_roots(tmp_path: Path):
    path = tmp_path / "notes.txt"
    path.write_text("hello")
    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    discovered = ingester.discover_files([MemoryRoot(scope="project:test", path=tmp_path)])
    assert discovered == [(MemoryRoot(scope="project:test", path=tmp_path), path)]


def test_discover_files_skips_virtualenv_and_cache_dirs(tmp_path: Path):
    kept = tmp_path / "docs" / "guide.md"
    skipped = tmp_path / ".venv" / "notes.md"
    kept.parent.mkdir(parents=True)
    skipped.parent.mkdir(parents=True)
    kept.write_text("keep me")
    skipped.write_text("skip me")
    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    discovered = ingester.discover_files([MemoryRoot(scope="project:test", path=tmp_path)])
    assert discovered == [(MemoryRoot(scope="project:test", path=tmp_path), kept)]


def test_discover_files_skips_generated_dirs_and_volume_state(tmp_path: Path):
    kept_result = tmp_path / ".claude" / "sprint-state" / "results" / "summary.md"
    skipped_volume = tmp_path / "docker" / "volume" / "state" / "ses" / "message.txt"
    skipped_target = tmp_path / "backend" / "target" / "surefire-reports" / "TEST-report.xml"
    skipped_dist = tmp_path / "ui" / "dist" / "assets" / "app.js"
    skipped_review = tmp_path / ".claude" / "sprint-state" / "reviews" / "note.md"
    skipped_playwright = tmp_path / "ui" / "playwright-report" / "index.html"
    skipped_results = tmp_path / "ui" / "test-results" / "results.json"

    for path in [
        kept_result,
        skipped_volume,
        skipped_target,
        skipped_dist,
        skipped_review,
        skipped_playwright,
        skipped_results,
    ]:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("content")

    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    discovered = ingester.discover_files([MemoryRoot(scope="project:test", path=tmp_path)])

    assert discovered == [(MemoryRoot(scope="project:test", path=tmp_path), kept_result)]


def test_file_extractor_rejects_unsupported_code_files(tmp_path: Path):
    path = tmp_path / "Example.kt"
    path.write_text("class Example {\n    fun greet() = \"hello\"\n}\n")
    extractor = FileExtractor()
    try:
        extractor.extract(path)
    except FileExtractionError as exc:
        assert "Unsupported file type" in str(exc)
    else:
        raise AssertionError("Expected code file extraction to be rejected")


def test_file_extractor_raises_file_extraction_error_for_invalid_pdf(tmp_path: Path):
    path = tmp_path / "invalid.pdf"
    path.write_bytes(b"AAAAA this is not a real pdf")
    extractor = FileExtractor()
    try:
        extractor.extract(path)
    except FileExtractionError as exc:
        assert "Failed to extract PDF" in str(exc)
    else:
        raise AssertionError("Expected invalid PDF extraction to fail")


import openpyxl


def test_file_extractor_extracts_xlsx_with_multiple_sheets(tmp_path: Path):
    path = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Revenue"
    ws1.append(["Quarter", "Revenue", "Growth"])
    ws1.append(["Q1 2025", 1200000, "15%"])
    ws1.append(["Q2 2025", 1380000, "15%"])
    ws2 = wb.create_sheet("Expenses")
    ws2.append(["Category", "Amount"])
    ws2.append(["Infra", 50000])
    wb.save(path)

    extractor = FileExtractor()
    result = extractor.extract(path)
    assert "## Sheet: Revenue" in result.text
    assert "Quarter | Revenue | Growth" in result.text
    assert "Q1 2025 | 1200000 | 15%" in result.text
    assert "## Sheet: Expenses" in result.text
    assert "Infra | 50000" in result.text


def test_file_extractor_skips_empty_xlsx_sheets(tmp_path: Path):
    path = tmp_path / "sparse.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Empty"
    ws2 = wb.create_sheet("HasData")
    ws2.append(["Name", "Value"])
    ws2.append(["Alpha", 1])
    wb.save(path)

    extractor = FileExtractor()
    result = extractor.extract(path)
    assert "Empty" not in result.text
    assert "## Sheet: HasData" in result.text


def test_file_extractor_extracts_xls():
    path = Path(__file__).parent / "fixtures" / "sample.xls"
    extractor = FileExtractor()
    result = extractor.extract(path)
    assert "## Sheet: Revenue" in result.text
    assert "Quarter | Revenue" in result.text
    assert "Q1 2025" in result.text
    assert "## Sheet: Expenses" in result.text
    assert "Infra" in result.text


def test_file_extractor_extracts_archimate():
    path = Path(__file__).parent / "fixtures" / "sample.archimate"
    extractor = FileExtractor()
    result = extractor.extract(path)
    assert "## Business Layer" in result.text
    assert "## Application Layer" in result.text
    assert "## Technology Layer" in result.text
    assert 'Business Process: "Order Fulfillment"' in result.text
    assert "Handles the end-to-end order lifecycle" in result.text
    assert 'Business Actor: "Customer"' in result.text
    assert 'Application Component: "Payment Gateway"' in result.text
    assert "## Relationships" in result.text
    assert "Order Fulfillment" in result.text
    assert "Customer" in result.text


def test_file_extractor_extracts_archimate_without_documentation(tmp_path: Path):
    path = tmp_path / "minimal.archimate"
    path.write_text("""<?xml version="1.0" encoding="UTF-8"?>
<archimate:model xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"
                 xmlns:archimate="http://www.archimatetool.com/archimate"
                 name="Minimal">
  <folder name="Business" type="business">
    <element xsi:type="archimate:BusinessProcess" name="Checkout" id="bp-1"/>
  </folder>
</archimate:model>""")
    extractor = FileExtractor()
    result = extractor.extract(path)
    assert 'Business Process: "Checkout"' in result.text


def test_discover_files_excludes_code_but_includes_documents(tmp_path: Path):
    py_file = tmp_path / "main.py"
    ts_file = tmp_path / "app.ts"
    md_file = tmp_path / "README.md"
    xlsx_file = tmp_path / "data.xlsx"

    for f in [py_file, ts_file, md_file]:
        f.write_text("content")
    wb = openpyxl.Workbook()
    wb.active.append(["test"])
    wb.save(xlsx_file)

    ingester = LocalFileIngester(FileExtractor(), DeterministicEmbedder(32))
    discovered = ingester.discover_files([MemoryRoot(scope="test", path=tmp_path)])
    discovered_paths = [p for _, p in discovered]
    assert md_file in discovered_paths
    assert xlsx_file in discovered_paths
    assert py_file not in discovered_paths
    assert ts_file not in discovered_paths
